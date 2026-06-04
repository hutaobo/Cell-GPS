from pathlib import Path
import re

from docx import Document
from docx.shared import Pt


HERE = Path(__file__).resolve().parent
INPUT = HERE / "manuscript_draft.md"
OUTPUT = HERE / "manuscript_draft.docx"


def clean_inline(text: str) -> str:
    text = text.replace("`", "")
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"\*(.*?)\*", r"\1", text)
    return text


def add_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.add_run(clean_inline(text))


def main() -> None:
    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)

    lines = INPUT.read_text(encoding="utf-8").splitlines()
    buffer: list[str] = []

    def flush() -> None:
        nonlocal buffer
        if buffer:
            add_paragraph(document, " ".join(buffer).strip())
            buffer = []

    for line in lines:
        stripped = line.strip()
        if not stripped:
            flush()
            continue
        if stripped.startswith("# "):
            flush()
            document.add_heading(clean_inline(stripped[2:]), level=0)
        elif stripped.startswith("## "):
            flush()
            document.add_heading(clean_inline(stripped[3:]), level=1)
        elif stripped.startswith("### "):
            flush()
            document.add_heading(clean_inline(stripped[4:]), level=2)
        elif stripped.startswith("- "):
            flush()
            document.add_paragraph(clean_inline(stripped[2:]), style="List Bullet")
        elif re.match(r"^\d+\. ", stripped):
            flush()
            document.add_paragraph(clean_inline(stripped), style="List Number")
        else:
            buffer.append(stripped)
    flush()
    document.save(OUTPUT)


if __name__ == "__main__":
    main()
