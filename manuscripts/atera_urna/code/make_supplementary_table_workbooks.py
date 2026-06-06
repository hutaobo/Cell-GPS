from pathlib import Path
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OUT = Path(__file__).resolve().parents[1]
TABLES = OUT / "tables"

HEADER_FILL = "D9EAF7"
README_FILL = "E8F4E8"
ALT_FILL = "F7F9FB"
FONT_NAME = "Arial"
BODY_SIZE = 10


SUMMARY_DESCRIPTIONS = {
    "total_parquet_rows": "Total rows in the source Atera transcript parquet file.",
    "gene_rows": "Rows corresponding to decoded gene transcripts.",
    "qv_ge_20_gene_rows": "Gene transcript rows retained after QV >= 20 quality filtering.",
    "qv_ge_20_unassigned_gene_rows": "High-quality gene transcripts with cell_id equal to UNASSIGNED.",
    "unassigned_fraction_of_qv_gene_rows": "Fraction of high-quality gene transcripts that were unassigned.",
    "genes_in_panel": "Number of assayed genes in the Atera WTA panel.",
    "genes_with_qv_ge_20_uRNA": "Genes with at least one high-quality unassigned RNA.",
    "genes_with_qv_ge_20_uRNA_fraction": "Fraction of assayed genes with at least one high-quality unassigned RNA.",
    "genes_analyzed_by_urna_only_COSTE": "Genes included in uRNA-COSTE and concordance analysis.",
    "median_spearman_vs_full": "Median Spearman concordance between uRNA-COSTE and all-transcript SSS profiles across analyzed genes.",
    "mean_spearman_vs_full": "Mean Spearman concordance between uRNA-COSTE and all-transcript SSS profiles across analyzed genes.",
    "q25_spearman_vs_full": "First quartile of Spearman concordance across analyzed genes.",
    "q75_spearman_vs_full": "Third quartile of Spearman concordance across analyzed genes.",
    "best_cluster_exact_match_rate": "Fraction of analyzed genes with identical uRNA-COSTE and all-transcript best cell groups.",
    "genes_spearman_lt_0_5": "Number of analyzed genes with Spearman concordance below 0.5.",
    "median_gene_uRNA_fraction_all_panel_genes": "Median gene-level uRNA fraction across the full assayed panel.",
    "median_gene_uRNA_fraction_analyzed_genes": "Median gene-level uRNA fraction across the 530 analyzed genes.",
}


TABLE_SPECS = [
    {
        "number": 1,
        "source": "summary_statistics.csv",
        "output": "supplementary_table_s1_summary_statistics.xlsx",
        "title": "Supplementary Table S1. Overall summary statistics",
        "sheet": "Table S1",
        "purpose": "Overall transcript, gene coverage and concordance statistics used in the manuscript.",
        "manuscript_use": "Supports the dataset-scale and global concordance statements in Results.",
    },
    {
        "number": 2,
        "source": "top_50_uRNA_genes.csv",
        "output": "supplementary_table_s2_top_50_urna_genes.xlsx",
        "title": "Supplementary Table S2. Highest-count uRNA genes",
        "sheet": "Table S2",
        "purpose": "The 50 genes with the highest high-quality uRNA counts in the full Atera panel.",
        "manuscript_use": "Supports the full-panel uRNA-coverage paragraph.",
    },
    {
        "number": 4,
        "source": "celltype_attribution_counts.csv",
        "output": "supplementary_table_s4_celltype_attribution_counts.xlsx",
        "title": "Supplementary Table S4. uRNA-COSTE best cell-group counts",
        "sheet": "Table S4",
        "purpose": "Counts of analyzed genes assigned to each uRNA-COSTE best cell group.",
        "manuscript_use": "Supports Figure 1d and the compartment-attribution paragraph.",
    },
    {
        "number": 6,
        "source": "marker_group_summary.csv",
        "output": "supplementary_table_s6_marker_group_summary.xlsx",
        "title": "Supplementary Table S6. Marker group summary",
        "sheet": "Table S6",
        "purpose": "Group-level marker-control statistics for curated marker classes.",
        "manuscript_use": "Supports the marker-control Results section.",
    },
    {
        "number": 7,
        "source": "marker_validation_table.csv",
        "output": "supplementary_table_s7_marker_validation.xlsx",
        "title": "Supplementary Table S7. Per-marker validation table",
        "sheet": "Table S7",
        "purpose": "Gene-level marker-control results for curated markers.",
        "manuscript_use": "Supports per-marker recovery and discordance statements.",
    },
    {
        "number": 8,
        "source": "lowest_concordance_genes.csv",
        "output": "supplementary_table_s8_lowest_concordance_genes.xlsx",
        "title": "Supplementary Table S8. Lowest-concordance genes",
        "sheet": "Table S8",
        "purpose": "The 30 genes with the lowest Spearman concordance between uRNA-COSTE and all-transcript SSS profiles.",
        "manuscript_use": "Supports Supplementary Fig. 1 and the discordant-profile Results section.",
    },
]


README_CONTENT = {
    3: {
        "title": "Supplementary Table S3. 530 analyzed genes for uRNA-COSTE and concordance analysis",
        "purpose": "Documents the exact 530-gene analysis universe used for uRNA-COSTE and concordance analyses.",
        "contents": "The Table S3 sheet lists gene selection source, marker status, transcript counts, uRNA fraction, uRNA-COSTE best cell group, all-transcript best cell group, match status and Spearman concordance.",
        "notes": "The 530-gene set is the union of the top 500 genes ranked by high-quality uRNA count and curated marker genes, with overlapping genes counted once.",
    },
    5: {
        "title": "Supplementary Table S5. Marker-control recovery in uRNA-COSTE",
        "purpose": "Reports gene-level curated marker-control results used to assess whether uRNAs recover interpretable tissue compartments.",
        "contents": "The Table S5 sheet lists marker group, expected compartment, assigned and unassigned transcript counts, uRNA-COSTE best cell group, all-transcript best cell group, match status and Spearman concordance for each marker gene.",
        "notes": "Group-level marker summaries are reported separately in Supplementary Table S6.",
    },
}


def clean_value(value):
    if pd.isna(value):
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return value


def read_table(spec):
    df = pd.read_csv(TABLES / spec["source"])
    if spec["number"] == 1:
        df["description"] = df["metric"].map(SUMMARY_DESCRIPTIONS)
    for col in df.columns:
        if df[col].dtype == bool:
            df[col] = df[col].map(lambda value: "TRUE" if value else "FALSE")
    return df


def header_fill_for(number):
    return README_FILL if number in {2, 5, 6} else HEADER_FILL


def style_sheet(ws, table_number):
    thin = Side(style="thin", color="D0D7DE")
    border = Border(top=thin, bottom=thin, left=thin, right=thin)
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    alt_fill = PatternFill("solid", fgColor=ALT_FILL)

    for cell in ws[1]:
        cell.font = Font(name=FONT_NAME, size=BODY_SIZE, bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        if row[0].row % 2 == 0:
            for cell in row:
                cell.fill = alt_fill
        for cell in row:
            cell.font = Font(name=FONT_NAME, size=BODY_SIZE)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for col_idx, cells in enumerate(ws.columns, start=1):
        max_len = 0
        for cell in cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        width = min(max(max_len + 2, 10), 58)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 24
    for row_idx in range(2, ws.max_row + 1):
        max_len = max(len("" if cell.value is None else str(cell.value)) for cell in ws[row_idx])
        ws.row_dimensions[row_idx].height = 18 if max_len < 80 else 32


def apply_number_formats(ws):
    percent_terms = ("fraction", "rate")
    spearman_terms = ("spearman", "rho")
    sss_terms = ("sss",)
    count_terms = ("count", "rows", "genes", "n_")

    headers = [cell.value for cell in ws[1]]
    for col_idx, header in enumerate(headers, start=1):
        name = str(header).lower()
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if not isinstance(cell.value, (int, float)):
                continue
            if any(term in name for term in percent_terms):
                cell.number_format = "0.0%"
            elif any(term in name for term in spearman_terms + sss_terms):
                cell.number_format = "0.000"
            elif any(term in name for term in count_terms):
                cell.number_format = "#,##0"
            else:
                cell.number_format = "0.000"

    if ws.title == "Table S3":
        for row_idx in range(2, ws.max_row + 1):
            metric = str(ws.cell(row=row_idx, column=1).value)
            value_cell = ws.cell(row=row_idx, column=2)
            if not isinstance(value_cell.value, (int, float)):
                continue
            if "fraction" in metric or "rate" in metric:
                value_cell.number_format = "0.0%"
            elif "spearman" in metric:
                value_cell.number_format = "0.000"
            elif "rows" in metric or "genes" in metric or "count" in metric:
                value_cell.number_format = "#,##0"
            else:
                value_cell.number_format = "0.000"


def add_readme(wb, spec):
    ws = wb.create_sheet("README")
    rows = readme_rows(spec["number"], spec)
    for row in rows:
        ws.append(row)

    fill = PatternFill("solid", fgColor=README_FILL)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(name=FONT_NAME, size=BODY_SIZE, bold=True)
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT_NAME, size=BODY_SIZE)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 115
    for row_idx in range(1, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 22 if row_idx == 1 else 32


def write_workbook(spec):
    df = read_table(spec)
    wb = Workbook()
    ws = wb.active
    ws.title = spec["sheet"]

    ws.append(list(df.columns))
    for _, row in df.iterrows():
        ws.append([clean_value(row[col]) for col in df.columns])

    style_sheet(ws, spec["number"])
    apply_number_formats(ws)
    add_readme(wb, spec)
    wb.save(TABLES / spec["output"])


def normalize_workbook(path):
    from openpyxl import load_workbook

    wb = load_workbook(path)
    standardize_sheet_names(wb, path)
    for ws in wb.worksheets:
        if ws.title == "README":
            normalize_readme_sheet(ws)
        else:
            style_sheet(ws, 0)
            apply_number_formats(ws)
    wb.save(path)


def table_number_from_filename(path):
    match = re.search(r"s(\d+)", path.name)
    return int(match.group(1)) if match else None


def standardize_sheet_names(wb, path):
    table_number = table_number_from_filename(path)
    if table_number is None:
        return

    primary = f"Table S{table_number}"
    rename_primary_sheet_if_needed(wb, primary)

    if table_number == 1:
        move_sheet_first(wb, primary)
    elif table_number == 2:
        if "Marker genes" in wb.sheetnames:
            wb["Marker genes"].title = primary
        move_sheet_first(wb, primary)
    else:
        move_sheet_first(wb, primary)

    remove_auxiliary_sheets(wb, primary)
    if "README" in wb.sheetnames:
        move_sheet_last(wb, "README")
        rewrite_readme_sheet(wb["README"], table_number)


def rename_primary_sheet_if_needed(wb, primary):
    if primary in wb.sheetnames:
        return
    for ws in wb.worksheets:
        if ws.title != "README":
            ws.title = primary
            return


def remove_auxiliary_sheets(wb, primary):
    keep = {primary, "README"}
    for ws in list(wb.worksheets):
        if ws.title not in keep:
            wb.remove(ws)


def move_sheet_first(wb, title):
    if title not in wb.sheetnames:
        return
    ws = wb[title]
    wb._sheets.remove(ws)
    wb._sheets.insert(0, ws)


def move_sheet_last(wb, title):
    if title not in wb.sheetnames:
        return
    ws = wb[title]
    wb._sheets.remove(ws)
    wb._sheets.append(ws)


def rewrite_readme_sheet(ws, table_number):
    rows = readme_rows(table_number)
    ws.delete_rows(1, ws.max_row)
    for row in rows:
        ws.append(row)


def readme_rows(table_number, spec=None):
    content = README_CONTENT.get(table_number)
    if content is None and spec is not None:
        content = {
            "title": spec["title"],
            "purpose": spec["purpose"],
            "contents": f"The Table S{table_number} sheet contains the data described by the Supplementary Table S{table_number} legend.",
            "notes": spec["manuscript_use"],
        }
    if content is None:
        content = {
            "title": f"Supplementary Table S{table_number}",
            "purpose": "Supplementary table accompanying the manuscript.",
            "contents": f"The Table S{table_number} sheet contains the primary table data.",
            "notes": "See the manuscript Supplementary Table legend for details.",
        }
    return [
        [content["title"], None],
        ["Purpose", content["purpose"]],
        ["Contents", content["contents"]],
        ["Notes", content["notes"]],
    ]


def normalize_readme_sheet(ws):
    fill = PatternFill("solid", fgColor=README_FILL)
    thin = Side(style="thin", color="D0D7DE")
    border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name=FONT_NAME, size=BODY_SIZE, bold=(cell.row == 1))
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = border
            if cell.row == 1:
                cell.fill = fill
    ws.freeze_panes = None
    ws.auto_filter.ref = None
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 115
    for row_idx in range(1, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 22 if row_idx == 1 else 32


def normalize_all_workbooks():
    for path in sorted(TABLES.glob("supplementary_table_s*.xlsx")):
        normalize_workbook(path)


def main():
    for spec in TABLE_SPECS:
        write_workbook(spec)
    normalize_all_workbooks()


if __name__ == "__main__":
    main()
