from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


HERE = Path(__file__).resolve().parents[1]
INPUT = HERE / "manuscript_draft.md"
OUTPUT = HERE / "manuscript_draft.docx"
FONT = "Times New Roman"
BODY_SIZE = Pt(12)
BLACK = RGBColor(0, 0, 0)
LINE_SPACING_RULE = WD_LINE_SPACING.ONE_POINT_FIVE
LINE_SPACING = 1.5
PARAGRAPH_AFTER = Pt(6)


def clean_inline(text: str) -> str:
    text = text.replace("`", "")
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"\*(.*?)\*", r"\1", text)
    text = re.sub(r"\^([A-Za-z0-9*#]+)", r"\1", text)
    return text


def set_run_font(run, size=BODY_SIZE, bold=False, italic=False) -> None:
    run.font.name = FONT
    run.font.size = size
    run.font.color.rgb = BLACK
    run.bold = bold
    run.italic = italic
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def set_style_font(style, size=BODY_SIZE, bold=False, italic=False) -> None:
    style.font.name = FONT
    style.font.size = size
    style.font.color.rgb = BLACK
    style.font.bold = bold
    style.font.italic = italic
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    force_style_text_black(style)
    remove_style_paragraph_borders(style)


def force_style_text_black(style) -> None:
    for color in style.element.findall(".//" + qn("w:color")):
        color.set(qn("w:val"), "000000")
        for attr in ("themeColor", "themeTint", "themeShade"):
            color.attrib.pop(qn(f"w:{attr}"), None)


def remove_style_paragraph_borders(style) -> None:
    ppr = style.element.find(qn("w:pPr"))
    if ppr is None:
        return
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)


def set_paragraph_format(
    paragraph,
    *,
    before=Pt(0),
    after=PARAGRAPH_AFTER,
    alignment=WD_ALIGN_PARAGRAPH.LEFT,
    first_line=None,
    left_indent=None,
) -> None:
    paragraph.alignment = alignment
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = LINE_SPACING_RULE
    fmt.line_spacing = LINE_SPACING
    fmt.space_before = before
    fmt.space_after = after
    if first_line is not None:
        fmt.first_line_indent = first_line
    if left_indent is not None:
        fmt.left_indent = left_indent


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = document.styles
    set_style_font(styles["Normal"])
    styles["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    styles["Normal"].paragraph_format.line_spacing_rule = LINE_SPACING_RULE
    styles["Normal"].paragraph_format.line_spacing = LINE_SPACING
    styles["Normal"].paragraph_format.space_before = Pt(0)
    styles["Normal"].paragraph_format.space_after = PARAGRAPH_AFTER

    set_style_font(styles["Title"], Pt(14), bold=True)
    styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Title"].paragraph_format.line_spacing_rule = LINE_SPACING_RULE
    styles["Title"].paragraph_format.line_spacing = LINE_SPACING
    styles["Title"].paragraph_format.space_before = Pt(0)
    styles["Title"].paragraph_format.space_after = Pt(12)

    set_style_font(styles["Heading 1"], BODY_SIZE, bold=True)
    styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    styles["Heading 1"].paragraph_format.line_spacing_rule = LINE_SPACING_RULE
    styles["Heading 1"].paragraph_format.line_spacing = LINE_SPACING
    styles["Heading 1"].paragraph_format.space_before = Pt(12)
    styles["Heading 1"].paragraph_format.space_after = PARAGRAPH_AFTER

    set_style_font(styles["Heading 2"], BODY_SIZE, bold=True, italic=True)
    styles["Heading 2"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    styles["Heading 2"].paragraph_format.line_spacing_rule = LINE_SPACING_RULE
    styles["Heading 2"].paragraph_format.line_spacing = LINE_SPACING
    styles["Heading 2"].paragraph_format.space_before = Pt(6)
    styles["Heading 2"].paragraph_format.space_after = PARAGRAPH_AFTER

    for list_style_name in ("List Bullet", "List Number"):
        set_style_font(styles[list_style_name])
        styles[list_style_name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        styles[list_style_name].paragraph_format.line_spacing_rule = LINE_SPACING_RULE
        styles[list_style_name].paragraph_format.line_spacing = LINE_SPACING
        styles[list_style_name].paragraph_format.space_before = Pt(0)
        styles[list_style_name].paragraph_format.space_after = PARAGRAPH_AFTER

    add_continuous_line_numbers(section)
    add_page_number(section)


def add_continuous_line_numbers(section) -> None:
    sect_pr = section._sectPr
    for child in list(sect_pr):
        if child.tag == qn("w:lnNumType"):
            sect_pr.remove(child)

    line_numbers = OxmlElement("w:lnNumType")
    line_numbers.set(qn("w:countBy"), "1")
    line_numbers.set(qn("w:restart"), "continuous")

    pg_mar = sect_pr.find(qn("w:pgMar"))
    if pg_mar is not None:
        sect_pr.insert(list(sect_pr).index(pg_mar) + 1, line_numbers)
    else:
        sect_pr.append(line_numbers)


def add_page_number(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    label = paragraph.add_run("Page ")
    set_run_font(label, Pt(10))

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")

    run = paragraph.add_run()
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(end)


def add_styled_run(
    paragraph,
    text: str,
    *,
    size=BODY_SIZE,
    bold=False,
    italic=False,
    superscript=False,
) -> None:
    if not text:
        return
    run = paragraph.add_run(text)
    set_run_font(run, Pt(9) if superscript else size, bold=bold, italic=italic)
    if superscript:
        run.font.superscript = True


def add_marked_runs(paragraph, text: str, *, size=BODY_SIZE, bold=False, italic=False) -> None:
    token_pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*|\^[A-Za-z0-9*#]+)")
    position = 0
    for match in token_pattern.finditer(text):
        if match.start() > position:
            add_styled_run(
                paragraph,
                text[position : match.start()],
                size=size,
                bold=bold,
                italic=italic,
            )

        token = match.group(0)
        if token.startswith("`"):
            add_styled_run(paragraph, token[1:-1], size=size, bold=bold, italic=italic)
        elif token.startswith("**"):
            add_styled_run(
                paragraph,
                token[2:-2],
                size=size,
                bold=True,
                italic=italic,
            )
        elif token.startswith("*"):
            add_styled_run(
                paragraph,
                token[1:-1],
                size=size,
                bold=bold,
                italic=True,
            )
        elif token.startswith("^"):
            add_styled_run(
                paragraph,
                token[1:],
                size=size,
                bold=bold,
                italic=italic,
                superscript=True,
            )
        position = match.end()

    if position < len(text):
        add_styled_run(paragraph, text[position:], size=size, bold=bold, italic=italic)


def add_text_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    set_paragraph_format(paragraph, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    add_marked_runs(paragraph, text)


def add_title(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Title")
    set_paragraph_format(paragraph, after=Pt(12), alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_marked_runs(paragraph, text, size=Pt(14), bold=True)


def add_heading(document: Document, text: str, level: int) -> None:
    style = "Heading 1" if level == 1 else "Heading 2"
    paragraph = document.add_paragraph(style=style)
    set_paragraph_format(
        paragraph,
        before=Pt(12 if level == 1 else 6),
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )
    add_marked_runs(paragraph, text, bold=True, italic=(level == 2))


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    set_paragraph_format(paragraph, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    add_marked_runs(paragraph, text)


def add_numbered(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Number")
    set_paragraph_format(paragraph, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    add_marked_runs(paragraph, re.sub(r"^\d+\.\s+", "", text))


def add_reference(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    set_paragraph_format(
        paragraph,
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        left_indent=Pt(18),
        first_line=Pt(-18),
    )
    add_marked_runs(paragraph, text)


def main() -> None:
    document = Document()
    configure_document(document)

    lines = INPUT.read_text(encoding="utf-8").splitlines()
    buffer: list[str] = []
    current_section = ""

    def flush() -> None:
        nonlocal buffer
        if buffer:
            add_text_paragraph(document, " ".join(buffer).strip())
            buffer = []

    for line in lines:
        stripped = line.strip()
        if not stripped:
            flush()
            continue
        if stripped.startswith("# "):
            flush()
            add_title(document, stripped[2:])
        elif stripped.startswith("## "):
            flush()
            current_section = clean_inline(stripped[3:])
            add_heading(document, current_section, level=1)
        elif stripped.startswith("### "):
            flush()
            add_heading(document, stripped[4:], level=2)
        elif stripped.startswith("- "):
            flush()
            add_bullet(document, stripped[2:])
        elif re.match(r"^\d+\. ", stripped):
            flush()
            if current_section == "References":
                add_reference(document, stripped)
            else:
                add_numbered(document, stripped)
        else:
            buffer.append(stripped)
    flush()
    document.save(OUTPUT)


if __name__ == "__main__":
    main()
